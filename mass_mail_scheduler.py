"""
Scheduled Mass Mail sender.

Companion to the existing MassEmail.exe tool, but instead of sending
immediately, each row in Schedule.xlsx carries its own "Date". The mail
for that row is sent SEND_DAYS_BEFORE days before that Date (currently 5),
independently of every other row. Date can also be a recurring monthly
spec like "10 Every month" -- see parse_recurring_day.

Example: Date = 15/09/2026, SEND_DAYS_BEFORE = 5 -> mail goes out 10/09/2026.

Run it once a day (see Run_Scheduler.bat / Windows Task Scheduler) and it
will send whatever has come due and leave the rest pending.

Sends over SMTP using the email id and Gmail App Password set with
set_sender_credentials.py / Set_Sender_Credentials.bat -- run that first.

Usage
    python mass_mail_scheduler.py                 check & send what's due
    python mass_mail_scheduler.py --dry-run        preview only, nothing sent
    python mass_mail_scheduler.py --force-today    send everything still
                                                    pending, ignoring dates
"""

from __future__ import annotations

import argparse
import calendar
import html
import json
import mimetypes
import os
import re
import smtplib
import sys
from datetime import date, datetime, timedelta
from email.message import EmailMessage
from pathlib import Path

import openpyxl

BASE = Path(__file__).resolve().parent
LEGAL_DIR = BASE.parent
DATA = LEGAL_DIR / "Data"

# Locally, attachments sit in ../Data/Attachments next to MassEmail.exe's own
# data. On GitHub Actions there is no such sibling folder -- the workflow
# sets ATTACHMENTS_DIR to point at the Attachments/ folder checked into the
# repo instead.
ATTACHMENTS = Path(os.environ["ATTACHMENTS_DIR"]) if os.environ.get("ATTACHMENTS_DIR") else DATA / "Attachments"

SCHEDULE = BASE / "Schedule.xlsx"
CREDENTIALS = BASE / "credentials.json"

EMAILS_SHEET = "Emails"

STATUS_COL = "Status"
SUBJECT_COL = "Subject"
MESSAGE_COL = "Message"

SEND_DAYS_BEFORE = 5  # mail always goes out this many days before the Date column

SMTP_HOST = "smtp.gmail.com"
SMTP_PORT = 465


class Problem(Exception):
    """Something the user has to fix before mail can go out."""


def _clean(value) -> str:
    if value is None:
        return ""
    return str(value).strip()


def load_credentials() -> dict:
    """LEGAL_MAILS_CREDENTIALS env var (GitHub Actions secret, JSON) takes priority;
    credentials.json (local runs, via set_sender_credentials.py) is the fallback."""
    env_creds = os.environ.get("LEGAL_MAILS_CREDENTIALS")
    if env_creds:
        return json.loads(env_creds)
    if CREDENTIALS.exists():
        return json.loads(CREDENTIALS.read_text())
    return {}


# --------------------------------------------------------------------------
# reading Schedule.xlsx
# --------------------------------------------------------------------------


def parse_date(value) -> date | None:
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value).strip()
    for fmt in ("%d/%m/%Y", "%d-%m-%Y", "%d.%m.%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return None


RECURRING_MONTHLY_RE = re.compile(r"every\s*month", re.IGNORECASE)
DAY_NUMBER_RE = re.compile(r"\d{1,2}")


def parse_recurring_day(value) -> int | None:
    """'10 Every month' / 'Every month on 10' -> 10. None if not a recurring spec."""
    text = _clean(value)
    if not text or not RECURRING_MONTHLY_RE.search(text):
        return None
    m = DAY_NUMBER_RE.search(text)
    if not m:
        return None
    day = int(m.group())
    return day if 1 <= day <= 31 else None


def this_months_occurrence(day: int, today: date) -> date:
    """Day N of today's month, clamped to that month's last day (e.g. 31 in Feb -> 28/29)."""
    last_day = calendar.monthrange(today.year, today.month)[1]
    return date(today.year, today.month, min(day, last_day))


SENT_STATUS_RE = re.compile(r"^sent\s*\((\d{2}/\d{2}/\d{4})\)$", re.IGNORECASE)


def parse_sent_status_date(status: str) -> date | None:
    m = SENT_STATUS_RE.match(status.strip())
    return parse_date(m.group(1)) if m else None


# --------------------------------------------------------------------------
# Gmail (SMTP + App Password)
# --------------------------------------------------------------------------


def get_smtp(from_email: str | None, app_password: str | None) -> smtplib.SMTP_SSL:
    if not from_email or not app_password:
        raise Problem(
            "No sender email/app password configured. Run Set_Sender_Credentials.bat "
            "(or `python set_sender_credentials.py`) first."
        )
    smtp = smtplib.SMTP_SSL(SMTP_HOST, SMTP_PORT)
    try:
        smtp.login(from_email, app_password)
    except smtplib.SMTPAuthenticationError as exc:
        smtp.quit()
        raise Problem(
            "Gmail rejected that email/app password. Make sure 2-Step Verification is "
            "turned on for the account and you're using an App Password (not the normal "
            "account password) -- generate one at https://myaccount.google.com/apppasswords"
        ) from exc
    return smtp


def to_html(message) -> str:
    text = _clean(message)
    if not text:
        return "<p></p>"
    if "<" in text and ">" in text:
        return text
    blocks = [f"<p>{html.escape(line)}</p>" for line in text.split("\n\n")]
    return "\n\n".join(blocks)


def build_message(row: dict, from_email: str) -> EmailMessage:
    to_addr = row["email"]
    cc_list = [c for c in (row.get("cc"), row.get("cc1"), row.get("cc2")) if c]
    subject = row.get(SUBJECT_COL) or ""
    body = row.get(MESSAGE_COL) or ""
    if not subject or not body:
        raise Problem(f"Missing {SUBJECT_COL if not subject else MESSAGE_COL} for this row.")

    msg = EmailMessage()
    msg["To"] = to_addr
    msg["From"] = from_email
    if cc_list:
        msg["Cc"] = ", ".join(cc_list)
    msg["Subject"] = subject
    msg.set_content("This message requires an HTML-capable mail client.")
    msg.add_alternative(to_html(body), subtype="html")

    for key in ("Attachment", "Attachment2"):
        name = row.get(key)
        if not name:
            continue
        path = ATTACHMENTS / name
        if not path.exists():
            raise Problem(f"Attachment not found: {path}")
        ctype, _ = mimetypes.guess_type(str(path))
        maintype, subtype = (ctype.split("/", 1) if ctype else ("application", "octet-stream"))
        msg.add_attachment(path.read_bytes(), maintype=maintype, subtype=subtype, filename=path.name)
    return msg


def send(smtp: smtplib.SMTP_SSL, msg: EmailMessage) -> None:
    smtp.send_message(msg)


# --------------------------------------------------------------------------
# main run
# --------------------------------------------------------------------------


def run(dry_run: bool, force_today: bool) -> None:
    if not SCHEDULE.exists():
        raise Problem(f"File not found: {SCHEDULE}. Fill in Schedule.xlsx first (see README.md).")

    wb = openpyxl.load_workbook(SCHEDULE)
    if EMAILS_SHEET not in wb.sheetnames:
        raise Problem(f"{SCHEDULE.name} has no '{EMAILS_SHEET}' sheet.")
    ws = wb[EMAILS_SHEET]

    headers = [_clean(ws.cell(1, c).value) for c in range(1, ws.max_column + 1)]
    header_col = {h: i + 1 for i, h in enumerate(headers) if h}
    if "Date" not in header_col:
        raise Problem(f"'{EMAILS_SHEET}' has no 'Date' column.")
    if "email" not in header_col:
        raise Problem(f"'{EMAILS_SHEET}' has no 'email' column.")

    changed = False
    for col in (STATUS_COL,):
        if col not in header_col:
            new_col = ws.max_column + 1
            ws.cell(1, new_col).value = col
            header_col[col] = new_col
            changed = True

    creds = load_credentials()
    from_email = creds.get("from_email")
    app_password = creds.get("app_password")
    print(f"Sending as: {from_email}" if from_email else
          "Sending as: (not configured -- run Set_Sender_Credentials.bat)")
    today = date.today()

    smtp = None if dry_run else get_smtp(from_email, app_password)

    due_count = 0
    sent_count = 0
    try:
        for r in range(2, ws.max_row + 1):
            get = lambda name: ws.cell(r, header_col[name]).value if name in header_col else None

            email = _clean(get("email"))
            if not email:
                continue

            status = _clean(get(STATUS_COL))
            date_value = get("Date")
            recurring_day = parse_recurring_day(date_value)

            if recurring_day:
                target = this_months_occurrence(recurring_day, today)
            else:
                target = parse_date(date_value)
                if not target:
                    ws.cell(r, header_col[STATUS_COL]).value = "Error: bad/missing Date"
                    changed = True
                    continue

            send_on = target - timedelta(days=SEND_DAYS_BEFORE)

            if status.lower().startswith("sent"):
                last_sent = parse_sent_status_date(status)
                already_sent_this_cycle = (
                    last_sent is not None
                    and recurring_day
                    and (last_sent.year, last_sent.month) == (send_on.year, send_on.month)
                )
                # A one-time row (or a recurring row whose sent date we can't parse) is
                # never resent. A recurring row only skips if already sent THIS month.
                if not recurring_day or already_sent_this_cycle or last_sent is None:
                    continue

            if not force_today and send_on > today:
                ws.cell(r, header_col[STATUS_COL]).value = (
                    f"Pending (sends {send_on:%d/%m/%Y}, {SEND_DAYS_BEFORE} day(s) before Date {target:%d/%m/%Y})"
                )
                changed = True
                continue

            due_count += 1
            row = {
                "email": email,
                "cc": _clean(get("cc")),
                "cc1": _clean(get("cc1")),
                "cc2": _clean(get("cc2")),
                "Attachment": _clean(get("Attachment")),
                "Attachment2": _clean(get("Attachment2")),
                SUBJECT_COL: _clean(get(SUBJECT_COL)),
                MESSAGE_COL: get(MESSAGE_COL),
            }

            tag = " [recurring monthly]" if recurring_day else ""
            try:
                msg = build_message(row, from_email or "(not configured)")
                if dry_run:
                    print(f"[DRY RUN] would send to {email}  (Date {target:%d/%m/%Y}, sends {send_on:%d/%m/%Y}){tag}")
                else:
                    send(smtp, msg)
                    ws.cell(r, header_col[STATUS_COL]).value = f"Sent ({today:%d/%m/%Y})"
                    changed = True
                    sent_count += 1
                    print(f"Sent to {email}  (Date {target:%d/%m/%Y}, sends {send_on:%d/%m/%Y}){tag}")
            except Problem as exc:
                ws.cell(r, header_col[STATUS_COL]).value = f"Error: {exc}"
                changed = True
                print(f"! {email}: {exc}", file=sys.stderr)
    finally:
        if smtp is not None:
            smtp.quit()

    if changed and not dry_run:
        wb.save(SCHEDULE)

    tail = " (dry run: nothing sent, nothing saved)" if dry_run else ""
    print(f"\n{due_count} row(s) due today, {sent_count} sent.{tail}")


def main() -> int:
    parser = argparse.ArgumentParser(description="Send scheduled mass-mail rows whose date has come due.")
    parser.add_argument("--dry-run", action="store_true", help="Preview only; sends nothing, saves nothing")
    parser.add_argument("--force-today", action="store_true",
                         help="Send everything still pending now, ignoring the Date column")
    args = parser.parse_args()

    try:
        run(dry_run=args.dry_run, force_today=args.force_today)
    except Problem as exc:
        print(f"\nError: {exc}", file=sys.stderr)
        return 1
    except KeyboardInterrupt:
        print("\nStopped.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
